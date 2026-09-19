"""
文档解析器测试(把各种格式的文件读成文字)
========================================
小白理解:用户上传的商品资料有 PDF、Word、Excel、CSV、纯文本五种格式。
这些测试就是"模拟用户上传文件",检查系统能不能把里面的文字正确读出来。

为什么这块特别重要?因为解析错了,后面检索和回答全是错的 ——
这是整条流水线的"第一道关口",脏数据进不去才对。

测试用的文件都是现做的临时小文件,用完就删,不碰用户的真实资料。
"""
from pathlib import Path

import docx
import openpyxl
import pytest

from app.rag.loader import load_document


# ==================== 纯文本 / Markdown ====================

class TestTextLoader:
    """纯文本和 Markdown:最基础也最常用的格式"""

    def test_txt_split_by_blank_line(self, tmp_path: Path):
        """按空行分段:两段之间空一行,应该切成 2 个文字块"""
        f = tmp_path / "商品介绍.txt"
        f.write_text("星辰X1手机简介\n\n电池容量5000mAh", encoding="utf-8")
        docs = load_document(f)
        assert len(docs) == 2
        assert "星辰X1手机简介" in docs[0].page_content

    def test_md_read_same_way(self, tmp_path: Path):
        """Markdown 走同一套逻辑,标题和正文都要能读出来"""
        f = tmp_path / "说明书.md"
        f.write_text("# 冰箱使用说明\n\n静置两小时后再通电", encoding="utf-8")
        docs = load_document(f)
        assert len(docs) == 2
        assert any("静置两小时" in d.page_content for d in docs)

    def test_gbk_encoding(self, tmp_path: Path):
        """GBK 编码的中文文件也要能读(老文件常见 GBK,读不了会变乱码)"""
        f = tmp_path / "gbk文件.txt"
        f.write_bytes("商品名称:星辰X1".encode("gbk"))
        docs = load_document(f)
        assert "星辰X1" in docs[0].page_content

    def test_empty_file_raises(self, tmp_path: Path):
        """空文件要报错,而不是悄悄当成"读成功但没内容"混进知识库"""
        f = tmp_path / "空的.txt"
        f.write_text("", encoding="utf-8")
        with pytest.raises(ValueError, match="没有读到任何文字"):
            load_document(f)

    def test_whitespace_only_file_raises(self, tmp_path: Path):
        """只有空格换行的文件同样算空文件"""
        f = tmp_path / "全是空格.txt"
        f.write_text("   \n\n  \t \n", encoding="utf-8")
        with pytest.raises(ValueError):
            load_document(f)


# ==================== CSV ====================

class TestCsvLoader:
    """CSV:每行一条商品记录,要拼成"字段名: 值"的句子"""

    def test_rows_become_sentences(self, tmp_path: Path):
        """每行拼成一句话,字段名和值都在(这样搜"电池容量"能命中)"""
        f = tmp_path / "商品表.csv"
        f.write_text(
            "商品名称,价格,电池容量\n星辰X1,2999,5000mAh\n",
            encoding="utf-8",
        )
        docs = load_document(f)
        assert len(docs) == 1  # 表头不算数据,只有 1 行商品
        text = docs[0].page_content
        assert "商品名称: 星辰X1" in text
        assert "电池容量: 5000mAh" in text

    def test_empty_cell_skipped(self, tmp_path: Path):
        """空单元格要跳过,不能拼出"保修期: "这种没意义的片段"""
        f = tmp_path / "有空单元格.csv"
        f.write_text("型号,保修期,备注\nS5 Pro,,无\n", encoding="utf-8")
        docs = load_document(f)
        assert "保修期" not in docs[0].page_content  # 空值不出现
        assert "S5 Pro" in docs[0].page_content

    def test_row_number_recorded(self, tmp_path: Path):
        """要记下行号(引用时能说明"来自第几行")"""
        f = tmp_path / "带行号.csv"
        f.write_text("型号,价格\nA,100\nB,200\n", encoding="utf-8")
        docs = load_document(f)
        assert docs[0].metadata["row"] == 2  # 表头占第 1 行,数据从第 2 行起
        assert docs[1].metadata["row"] == 3

    def test_headerless_row_gets_column_name(self, tmp_path: Path):
        """数据列比表头多时,用"第N列"兜底,不能报错崩溃"""
        f = tmp_path / "列数不齐.csv"
        f.write_text("型号,价格\nA,100,额外的一列\n", encoding="utf-8")
        docs = load_document(f)
        assert "第3列: 额外的一列" in docs[0].page_content


# ==================== Excel ====================

class TestExcelLoader:
    """Excel:电商商品参数表最常见的格式"""

    def test_excel_rows_to_sentences(self, tmp_path: Path):
        """Excel 每行拼成一句话,和 CSV 逻辑一致"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["商品名称", "电池容量", "价格"])
        ws.append(["星辰X1", "5000mAh", 2999])
        f = tmp_path / "商品参数.xlsx"
        wb.save(f)

        docs = load_document(f)
        assert len(docs) == 1
        text = docs[0].page_content
        assert "商品名称: 星辰X1" in text
        assert "5000mAh" in text

    def test_number_converted_to_text(self, tmp_path: Path):
        """Excel 里的数字要能变成文字(否则拼句子时会报类型错误)"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["重量"])
        ws.append([0.185])
        f = tmp_path / "重量.xlsx"
        wb.save(f)

        docs = load_document(f)
        assert "0.185" in docs[0].page_content

    def test_multiple_sheets(self, tmp_path: Path):
        """多个工作表都要读(商品表和FAQ表常常分开放在两个 sheet)"""
        wb = openpyxl.Workbook()
        wb.active.title = "商品表"
        wb.active.append(["型号"])
        wb.active.append(["星辰X1"])
        ws2 = wb.create_sheet("FAQ")
        ws2.append(["问题", "答案"])
        ws2.append(["怎么退货", "七天无理由"])
        f = tmp_path / "多表.xlsx"
        wb.save(f)

        docs = load_document(f)
        assert len(docs) == 2
        assert {d.metadata["sheet"] for d in docs} == {"商品表", "FAQ"}

    def test_empty_row_skipped(self, tmp_path: Path):
        """整行空白要跳过,不能产生空片段污染知识库"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["型号", "价格"])
        ws.append([None, None])
        ws.append(["星辰X1", 2999])
        f = tmp_path / "有空白行.xlsx"
        wb.save(f)

        docs = load_document(f)
        assert len(docs) == 1


# ==================== Word ====================

class TestDocxLoader:
    """Word:说明书常见格式,正文和表格都要读"""

    def test_paragraphs(self, tmp_path: Path):
        """正文段落逐段读出"""
        d = docx.Document()
        d.add_paragraph("星辰X1手机使用说明")
        d.add_paragraph("首次使用请充满电")
        d.add_paragraph("")  # 空段落要跳过
        f = tmp_path / "说明书.docx"
        d.save(f)

        docs = load_document(f)
        assert len(docs) == 2
        assert docs[0].metadata["type"] == "段落"

    def test_table_becomes_field_value(self, tmp_path: Path):
        """表格数据行拼成"字段名: 值"(参数表是 Word 里最常见的表格)

        注意:这里第一行必须是字段名(型号/容量/重量),系统才认得出每个值的含义。
        """
        d = docx.Document()
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "电池容量"
        table.cell(0, 1).text = "重量"
        table.cell(1, 0).text = "5000mAh"
        table.cell(1, 1).text = "185g"
        f = tmp_path / "参数表.docx"
        d.save(f)

        docs = load_document(f)
        table_doc = next(x for x in docs if x.metadata["type"] == "表格" and "5000mAh" in x.page_content)
        assert "电池容量: 5000mAh" in table_doc.page_content
        assert "重量: 185g" in table_doc.page_content

    def test_header_row_defines_field_names(self, tmp_path: Path):
        """第一行被当作"字段名",后续行才拼成"字段名: 值"

        小白理解:表格像 Excel,第一行是表头。如果表头写的是"参数/数值"，
        系统就只能拼出"参数: 电池容量"，读起来别扭 —— 所以表头要写真实字段名。
        """
        d = docx.Document()
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "参数"
        table.cell(0, 1).text = "数值"
        table.cell(1, 0).text = "电池容量"
        table.cell(1, 1).text = "5000mAh"
        f = tmp_path / "表头不规范.docx"
        d.save(f)

        docs = load_document(f)
        # 表头原样保留一份(方便搜到字段名),数据行用表头做前缀
        assert any(x.page_content == "参数 | 数值" for x in docs)
        assert any("参数: 电池容量; 数值: 5000mAh" in x.page_content for x in docs)

    def test_header_row_kept(self, tmp_path: Path):
        """表头本身也存一份,保证"电池容量"这个词能被搜到"""
        d = docx.Document()
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "电池容量"
        table.cell(0, 1).text = "价格"
        table.cell(1, 0).text = "5000mAh"
        table.cell(1, 1).text = "2999"
        f = tmp_path / "表头.docx"
        d.save(f)

        docs = load_document(f)
        assert any("电池容量 | 价格" in x.page_content for x in docs)


# ==================== PDF ====================

class TestPdfLoader:
    """PDF:逐页读取并记页码,引用时能显示"来自第 3 页" """

    def test_pdf_pages(self, tmp_path: Path):
        """两页的 PDF 应该读出 2 个文字块,页码从 1 开始"""
        pymupdf = pytest.importorskip("pymupdf", reason="未安装 PyMuPDF")
        f = tmp_path / "说明书.pdf"

        doc = pymupdf.open()
        font = "china-s"  # PyMuPDF 内置的中文字体别名
        for page_text in ("第一页内容:星辰X1手机", "第二页内容:保修两年"):
            page = doc.new_page()
            page.insert_text((72, 100), page_text, fontname=font, fontsize=12)
        doc.save(f)
        doc.close()

        docs = load_document(f)
        assert len(docs) == 2
        assert docs[0].metadata["page"] == 1
        assert docs[1].metadata["page"] == 2
        assert "星辰X1" in docs[0].page_content


# ==================== 格式分发 ====================

class TestLoadDocument:
    """解析入口:根据文件后缀自动挑解析方式"""

    def test_unsupported_extension(self, tmp_path: Path):
        """不支持的格式要明确报错(而不是读出一堆乱码)"""
        f = tmp_path / "压缩包.zip"
        f.write_bytes(b"PK\x03\x04")
        with pytest.raises(ValueError, match="不支持的文件格式"):
            load_document(f)

    def test_extension_case_insensitive(self, tmp_path: Path):
        """后缀大写(.TXT)也要能认出来(用户重命名文件名很常见)"""
        f = tmp_path / "大写后缀.TXT"
        f.write_text("商品介绍", encoding="utf-8")
        docs = load_document(f)
        assert len(docs) == 1
