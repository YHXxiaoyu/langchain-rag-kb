"""
文档解析器(把各种格式的文件"读成文字")
=====================================
小白理解:用户上传的商品资料五花八门 —— PDF 说明书、Word 文档、Excel 参数表、
纯文本。这个文件就是"翻译官",负责把它们统统读成纯文字,后面的流程才能处理。

支持的格式与本项目的针对性设计:
  PDF  —— 逐页读取,记录页码(引用时可以显示"来自第 3 页")
  Word —— 段落 + 表格都要读(商品参数经常放在表格里)
  Excel—— 每行一个商品,自动拼成"字段名: 值"的句子(最适合商品参数表)
  CSV  —— 同 Excel
  TXT/Markdown —— 直接读
"""
from pathlib import Path

import docx
import openpyxl
import pymupdf  # PyMuPDF:读取 PDF 的库(注意不是老的 fitz 名字)
from langchain_core.documents import Document

from app.utils.logger import logger


def _load_pdf(path: Path) -> list[Document]:
    """解析 PDF:一页一个文档块,记录页码"""
    docs: list[Document] = []
    with pymupdf.open(path) as pdf:
        for page_no, page in enumerate(pdf, start=1):
            text = page.get_text().strip()
            if text:
                docs.append(Document(page_content=text, metadata={"page": page_no}))
    return docs


def _load_docx(path: Path) -> list[Document]:
    """解析 Word:正文段落逐段读取,表格逐行拼成句子"""
    d = docx.Document(str(path))
    docs: list[Document] = []

    # 1. 正文段落
    for para in d.paragraphs:
        text = para.text.strip()
        if text:
            docs.append(Document(page_content=text, metadata={"type": "段落"}))

    # 2. 表格(商品参数表常常是 Word 表格)
    for table_no, table in enumerate(d.tables, start=1):
        for row_no, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells]
            # 过滤掉整行都是空的
            if not any(cells):
                continue
            # 首行通常是表头,用它给每个单元格配上字段名
            header = [c.text.strip() for c in table.rows[0].cells] if len(table.rows) > 0 else []
            if row_no == 1:
                # 表头本身也存一份,保证"字段名"这个词能被搜到
                line = " | ".join(c for c in cells if c)
            else:
                # 数据行:拼成 "字段名: 值; 字段名: 值" 的形式
                parts = []
                for i, cell in enumerate(cells):
                    if not cell:
                        continue
                    name = header[i] if i < len(header) and header[i] else f"第{i + 1}列"
                    parts.append(f"{name}: {cell}")
                line = "; ".join(parts)
            if line.strip():
                docs.append(Document(page_content=line, metadata={"type": "表格", "table": table_no}))

    return docs


def _load_excel(path: Path) -> list[Document]:
    """
    解析 Excel:每行一个商品,拼成"字段名: 值"的句子。
    例如一行商品数据会变成:
      "商品名称: 星辰X1手机; 价格: 2999元; 电池容量: 5000mAh"
    这样检索"星辰X1 电池"时能精确命中。
    """
    docs: list[Document] = []
    wb = openpyxl.load_workbook(path, data_only=True)

    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        # 假设第一行是表头(字段名)
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]

        for row_no, row in enumerate(rows[1:], start=2):
            parts = []
            for col_no, value in enumerate(row):
                if value is None or str(value).strip() == "":
                    continue
                name = headers[col_no] if col_no < len(headers) and headers[col_no] else f"第{col_no + 1}列"
                parts.append(f"{name}: {str(value).strip()}")
            if parts:
                line = "; ".join(parts)
                docs.append(
                    Document(
                        page_content=line,
                        metadata={"type": "表格行", "sheet": sheet.title, "row": row_no},
                    )
                )
    return docs


def _load_csv(path: Path) -> list[Document]:
    """解析 CSV:逻辑同 Excel(每行一条记录)"""
    import csv

    docs: list[Document] = []
    # 中文 CSV 可能是 GBK 编码,先试 UTF-8 再退回 GBK
    for encoding in ("utf-8-sig", "gbk", "utf-8"):
        try:
            with open(path, encoding=encoding, newline="") as f:
                reader = csv.reader(f)
                rows = list(reader)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("CSV 文件编码无法识别(尝试了 UTF-8 和 GBK)")

    if not rows:
        return docs

    headers = [h.strip() for h in rows[0]]
    for row_no, row in enumerate(rows[1:], start=2):
        parts = []
        for col_no, value in enumerate(row):
            if not value.strip():
                continue
            name = headers[col_no] if col_no < len(headers) and headers[col_no] else f"第{col_no + 1}列"
            parts.append(f"{name}: {value.strip()}")
        if parts:
            docs.append(Document(page_content="; ".join(parts), metadata={"type": "表格行", "row": row_no}))
    return docs


def _load_text(path: Path) -> list[Document]:
    """解析纯文本 / Markdown:整体读入,按空行分段"""
    content = None
    # 中文文本常见编码:UTF-8 和 GBK,依次尝试
    for encoding in ("utf-8", "gbk", "utf-8-sig"):
        try:
            content = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        raise ValueError("文本文件编码无法识别(尝试了 UTF-8 和 GBK)")

    docs = []
    for block in content.split("\n\n"):
        block = block.strip()
        if block:
            docs.append(Document(page_content=block, metadata={"type": "文本"}))
    return docs


# 扩展名 → 对应的解析函数
_PARSERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".xlsx": _load_excel,
    ".csv": _load_csv,
    ".txt": _load_text,
    ".md": _load_text,
}


def load_document(file_path: Path) -> list[Document]:
    """
    解析入口:根据文件后缀自动选择合适的解析方式。
    返回一组"文字块"(每个块带页码/行号等来源信息,用于日后引用展示)。
    """
    ext = file_path.suffix.lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"不支持的文件格式: {ext}")

    docs = parser(file_path)
    logger.info(f"📄 已解析 {file_path.name}: 得到 {len(docs)} 个文字块")

    if not docs:
        raise ValueError("文件里没有读到任何文字(可能是扫描版 PDF 或空文件)")
    return docs
